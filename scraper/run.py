"""
Subsidie-scraper voor isolatieregelingen van ALLE Nederlandse gemeenten.

Per gemeente:
  1. Snel zoeken in het CVDR (lokaleregelgeving.overheid.nl).
  2. Gevonden? Het taalmodel leest de voorwaarden uit (bedrag, WOZ, bouwjaar, ...).
     Het gratis quotum is niet genoeg voor alles in één keer: wat niet lukt,
     volgt bij de volgende run. Al uitgelezen regelingen worden alleen opnieuw
     gelezen als de CVDR-versie verandert.
  3. Voor ELKE gemeente (eens per web_zoeken_elke_dagen): zoeken op de gemeentesite
     (sitemap) en bij partners (energieloketten, bouwloketten), links naar voorwaarden en
     pdf's volgen, en als dat niets oplevert een AI-zoekactie met Google (webbronnen.py).
     Alle pagina's van een gemeente gaan samen in één AI-aanroep, alleen als de tekst veranderd is.
     Zulke regelingen krijgen "bron": "web": altijd zelf controleren.
  4. Bij elke uitgelezen regeling wordt gecontroleerd of het "bewijs" van het model echt in de
     brontekst staat. Zo niet: lagere betrouwbaarheid en een melding in de PR.

Schrijft regelingen.json (de checkpagina) en zoekstatus.json (per gemeente:
wat nog wacht op uitlezen en wanneer er op internet is gezocht).

Draait in GitHub Actions. Taalmodel kies je in scraper/config.json ("provider":
"gemini" of "claude"). Lokaal testen: pip install requests google-genai anthropic
  python scraper/run.py                       normale run
  python scraper/run.py --forceer             alles opnieuw extraheren
  python scraper/run.py --alleen Doesburg,Arnhem   alleen deze gemeenten
  python scraper/run.py --vergelijk tests/baseline_pilot.json
  python scraper/run.py --alleen-overzicht    alleen overzicht_voorwaarden.xlsx opnieuw maken
  python scraper/run.py --zonder-ai --alleen Zeist   alleen zoeken (CVDR, site, partners), geen taalmodel
"""
import datetime as dt
import hashlib
import html
import json
import os
import pathlib
import re
import signal
import sys
import time

import requests

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from normaliseer import normaliseer  # noqa: E402  (getallen als getal, ja/nee als true/false)
import export_tios  # noqa: E402  (tios/subsidies.json en .csv voor de koppeling met TIOS)
import webbronnen  # noqa: E402  (gemeentesite, partners en AI-zoeken)

ROOT = pathlib.Path(__file__).resolve().parent.parent
CFG = json.loads((ROOT / "scraper/config.json").read_text(encoding="utf-8"))
PROMPT = (ROOT / "scraper/prompt_extractie.md").read_text(encoding="utf-8")
OUT = ROOT / "regelingen.json"
DATA = ROOT / "data"
STATE = DATA / "state.json"
DEKKING = DATA / "dekking.json"
LOG_ALL = DATA / "wijzigingen.md"
LOG_LAST = DATA / "wijzigingen_laatste.md"
ALLE_GEMEENTEN = DATA / "alle_gemeenten.json"  # buffer voor als PDOK even niet reageert
ZOEKSTATUS = ROOT / "zoekstatus.json"  # per gemeente: wachtrij + laatste zoekactie op internet
WEB_STATE = DATA / "web_state.json"  # wanneer per gemeente op internet is gezocht
OVERZICHT = ROOT / "overzicht_voorwaarden.xlsx"  # alles in één sheet, om te lezen
VOORTGANG = DATA / "voortgang.json"  # bij welke gemeente de volgende run begint
GEMEENTE_SITES = DATA / "gemeente_sites.json"  # webadres per gemeente (uit het overheidsregister)
SRU = "https://zoekservice.overheid.nl/sru/Search"
UA = {"User-Agent": "Takkenkamp-subsidiecheck/0.1 (interne tool)"}
VANDAAG = dt.date.today().isoformat()
DATA.mkdir(parents=True, exist_ok=True)  # map data/ aanmaken als die ontbreekt
FORCEER = "--forceer" in sys.argv
PROVIDER = CFG.get("provider", "gemini")
MODELLEN = CFG["modellen"][PROVIDER]
PAUZE = CFG.get("pauze_tussen_aanroepen_sec", 7)  # gratis Gemini: max. enkele verzoeken per minuut
START = time.monotonic()
MAX_AI_SEC = CFG.get("max_minuten_taalmodel", 100) * 60  # daarna: rest bij de volgende run
MAX_RUN_SEC = CFG.get("max_minuten_run", 300) * 60  # daarna netjes stoppen, ruim vóór de grens van GitHub
BEWAAR_ELKE_SEC = 300  # tussentijds opslaan, zodat een harde onderbreking weinig werk kost
ALLEEN = [n.strip().lower() for n in sys.argv[sys.argv.index("--alleen") + 1].split(",")] if "--alleen" in sys.argv else None
ZONDER_AI = "--zonder-ai" in sys.argv  # alleen zoeken, niets naar het taalmodel (om het zoeken te testen)
WEB_VERSIE = 2  # omhoog als het zoeken op internet verandert: dan worden alle gemeenten opnieuw doorzocht
FILTER_VERSIE = 2  # omhoog als de filters ruimer worden: eerder afgewezen regelingen worden opnieuw bekeken

# --aanbieder github: de hoofdaanbieder overslaan en meteen bij deze reserve-aanbieder beginnen (om te testen)
START_AANBIEDER = sys.argv[sys.argv.index("--aanbieder") + 1].strip().lower() if "--aanbieder" in sys.argv else None
if START_AANBIEDER in ("", "standaard", PROVIDER):
    START_AANBIEDER = None
HOOFD_SLEUTEL = "GEMINI_API_KEY" if PROVIDER == "gemini" else "ANTHROPIC_API_KEY"

if "--alleen-overzicht" in sys.argv or ZONDER_AI or START_AANBIEDER or not os.environ.get(HOOFD_SLEUTEL):
    client = None  # geen hoofdaanbieder: geen sleutel, of alleen reserve-aanbieders gebruiken
elif PROVIDER == "gemini":
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"],
                          http_options=types.HttpOptions(timeout=180_000))  # max. 3 min per aanroep
else:
    from anthropic import Anthropic
    client = Anthropic()  # leest ANTHROPIC_API_KEY


# ---------- hulpfuncties ----------
def heeft_voorwaarden(r):
    """Staat er minstens één toetsbare voorwaarde in de regeling?"""
    return any(v not in (None, "", [], {}) for k, v in (r.get("criteria") or {}).items()
               if not re.search(r"opmerking|peildatum|regel$|uitzondering", k))


def schrijf_samenvatting(resultaten, stand, totaal):
    """Overzicht per gemeente: in de log en op de samenvattingspagina van de GitHub-run."""
    volgorde = {"❌": 0, "🛑": 1, "⚠️": 2, "🔎": 3, "➖": 4, "✅": 5}
    tel = {}
    for _, sym, _ in resultaten:
        tel[sym] = tel.get(sym, 0) + 1
    kop = (f"{len(resultaten)} van {totaal} gemeenten bekeken: "
           + ", ".join(f"{sym} {tel[sym]}" for sym in sorted(tel, key=lambda x: volgorde.get(x, 9))))
    print("\n== Samenvatting ==\n" + kop)
    print("✅ voorwaarden opgehaald · ⚠️ deels gelukt of zonder voorwaarden · ➖ geen (open) regeling · "
          "❌ mislukt · 🛑 gestopt (budget/tijd)")
    pad = os.environ.get("GITHUB_STEP_SUMMARY")
    if not pad:
        return
    regels = ["## Subsidie-update: resultaat per gemeente", "", kop, ""]
    if stand.get("volgende"):
        regels += [f"**Gestopt ({stand['reden']}).** De volgende run gaat verder bij **{stand['volgende']}**.", ""]
    regels += ["✅ voorwaarden opgehaald · ⚠️ deels gelukt of zonder voorwaarden · ➖ geen (open) regeling · "
               "❌ mislukt · 🛑 gestopt", "", "| | Gemeente | Resultaat |", "|---|---|---|"]
    for naam, sym, tekst in sorted(resultaten, key=lambda x: (volgorde.get(x[1], 9), x[0])):
        regels.append(f"| {sym} | {naam} | {tekst.replace('|', '/')} |")
    with open(pad, "a", encoding="utf-8") as f:
        f.write("\n".join(regels) + "\n")


def limiet_reden(fout):
    """Korte reden voor de PR-tekst waarom de run stopte."""
    return "AI-tijd van deze run op" if "tijdslimiet" in str(fout) else "AI-budget of daglimiet op"


def lees(p, standaard):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else standaard


def schrijf(p, obj):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


class LimietOp(Exception):
    """De (gratis) limiet van het taalmodel is bereikt; de rest volgt bij de volgende run."""


TOKENS = {"extractie": 16000, "zoeken": 4000, "filter": 1000}  # web-extractie gebruikt "extractie"
UITGEPUT = set()  # modellen waarvan de daglimiet op is: deze run niet meer proberen


def is_limiet(fout):
    """429 = limiet, 402 = tegoed op. Beide: stoppen of reservemodel, niet blijven proberen."""
    return re.search(r"\b(429|402)\b|RESOURCE_EXHAUSTED|rate_limit|credits", fout) is not None


def is_daglimiet(fout):
    return re.search(r"\b402\b|credits|per ?day", fout, re.I) is not None


class ZonderAI(Exception):
    """Run met --zonder-ai: het taalmodel wordt niet gebruikt."""


def bronnen_gemini(r):
    """De webadressen die Google bij een zoekactie echt heeft gebruikt (doorverwijzing; volgt bij ophalen)."""
    uit = []
    for c in getattr(r, "candidates", None) or []:
        gm = getattr(c, "grounding_metadata", None)
        for ch in (getattr(gm, "grounding_chunks", None) or []) if gm else []:
            web = getattr(ch, "web", None)
            if web is not None and getattr(web, "uri", None):
                uit.append(web.uri)
    return uit


def bronnen_claude(r):
    uit = []
    for b in r.content:
        if getattr(b, "type", "") == "web_search_tool_result" and isinstance(getattr(b, "content", None), list):
            uit += [x.url for x in b.content if getattr(x, "url", None)]
    return uit


def reserve_aanbieders():
    """Reserve-aanbieders uit config.json met een sleutel (GitHub Models gebruikt het token van de workflow)."""
    lijst = [a for a in CFG.get("reserve_aanbieders", []) if os.environ.get(a.get("sleutel_env", ""))]
    if START_AANBIEDER:
        namen = [a["naam"].lower() for a in lijst]
        lijst = lijst[namen.index(START_AANBIEDER):] if START_AANBIEDER in namen else []
    return lijst


class TeLang(Exception):
    """De tekst past niet in het model van deze aanbieder."""


def llm_openai(a, taak, tekst, json_uit):
    """Eén aanroep bij een OpenAI-compatibele aanbieder (GitHub Models, Cerebras, Groq, Mistral, OpenRouter).
    Geeft de tekst, of None als de aanbieder op is of weigert (dan de volgende proberen)."""
    model = (a.get("modellen") or {}).get(taak)
    if not model or a["naam"] in UITGEPUT:
        return None
    if len(tekst) > a.get("max_invoer_tekens", 10 ** 9):
        raise TeLang(f"{a['naam']}: tekst te lang ({len(tekst)} tekens, max. {a['max_invoer_tekens']})")
    body = {"model": model, "messages": [{"role": "user", "content": tekst}], "temperature": 0,
            "max_tokens": min(TOKENS[taak], a.get("max_uitvoer_tokens", 4000))}
    if json_uit and a.get("json", True):
        body["response_format"] = {"type": "json_object"}
    kop = {"Authorization": f"Bearer {os.environ[a['sleutel_env']]}", "Content-Type": "application/json"}
    for poging in range(3):
        time.sleep(a.get("pauze_sec", 2))
        try:
            r = requests.post(a["basis_url"].rstrip("/") + "/chat/completions", headers=kop, json=body, timeout=180)
        except requests.RequestException as e:
            print(f"  {a['naam']}: fout ({str(e)[:100]}), poging {poging + 1}/3")
            time.sleep(10)
            continue
        if r.status_code == 200:
            try:
                antw = r.json()["choices"][0]["message"]["content"] or ""
            except (ValueError, KeyError, IndexError):
                antw = ""
            if antw.strip():
                print(f"  (antwoord van {a['naam']}, {model})")
                return antw
            continue
        fout = r.text[:200].replace("\n", " ")
        if r.status_code == 429:
            wacht = int(re.sub(r"\D", "", r.headers.get("retry-after", "")) or 0)
            if 0 < wacht <= 65 and poging < 2:
                time.sleep(wacht)
                continue
            UITGEPUT.add(a["naam"])
            print(f"  {a['naam']}: limiet bereikt ({fout[:100]}), deze run niet meer gebruiken")
            return None
        if r.status_code in (401, 403):
            UITGEPUT.add(a["naam"])
            print(f"  {a['naam']}: geen toegang ({r.status_code}: {fout[:100]})")
            return None
        if r.status_code in (400, 413) and re.search(r"token|context|too (large|long)|length", fout, re.I):
            raise TeLang(f"{a['naam']}: {fout[:120]}")
        if r.status_code == 400 and "response_format" in body:
            body.pop("response_format")   # niet elke aanbieder kent de JSON-modus
            continue
        print(f"  {a['naam']}: fout {r.status_code} ({fout[:100]}), poging {poging + 1}/3")
        time.sleep(10 * (poging + 1))
    return None


def llm(taak, tekst, json_uit=False, zoeken=False, met_bronnen=False):
    """Stuurt tekst naar het taalmodel. Eerst de hoofdaanbieder (config "provider"); is die op of er is
    geen sleutel, dan de reserve-aanbieders uit config.json, op volgorde. Alles op: LimietOp.
    zoeken=True: het model mag op internet zoeken (alleen Gemini en Claude kunnen dat).
    met_bronnen=True: geeft (tekst, [webadressen die de zoekmachine echt gaf])."""
    if ZONDER_AI:
        raise ZonderAI()
    if time.monotonic() - START > MAX_AI_SEC:
        raise LimietOp("tijdslimiet van deze run bereikt")
    fout = None
    if client is not None:
        try:
            return llm_hoofd(taak, tekst, json_uit, zoeken, met_bronnen)
        except LimietOp as e:
            fout = e
    if zoeken:
        raise LimietOp(f"zoeken op internet kan alleen met {PROVIDER}" + (f" ({fout})" if fout else ""))
    te_lang = []
    for a in reserve_aanbieders():
        try:
            antw = llm_openai(a, taak, tekst, json_uit)
        except TeLang as e:
            print(f"  {e}")
            te_lang.append(a["naam"])
            continue
        if antw is not None:
            return (antw, []) if met_bronnen else antw
    if te_lang:   # niet op, maar deze tekst is te lang voor wat er nog over is: later opnieuw
        raise RuntimeError(f"tekst te lang voor de beschikbare modellen ({', '.join(te_lang)})")
    raise LimietOp(str(fout or "geen taalmodel beschikbaar (sleutel ontbreekt of alle limieten op)")[:200])


def llm_hoofd(taak, tekst, json_uit=False, zoeken=False, met_bronnen=False):
    """De hoofdaanbieder (Gemini of Claude). Bij drukte of limiet: eerst even wachten,
    dan het reservemodel proberen, en anders LimietOp opgooien."""
    modellen = [MODELLEN[taak]] + ([MODELLEN[taak + "_reserve"]] if MODELLEN.get(taak + "_reserve") else [])
    modellen = [m for m in modellen if m not in UITGEPUT]
    if not modellen:
        raise LimietOp(f"daglimiet bereikt voor {taak}")
    laatste_fout = None
    for model in modellen:
        for poging in range(3):
            try:
                time.sleep(PAUZE)
                if PROVIDER == "gemini":
                    cfg = types.GenerateContentConfig(
                        temperature=0,
                        max_output_tokens=TOKENS[taak],
                        # JSON-modus gaat niet samen met Google zoeken; dan zelf de JSON uit de tekst halen
                        response_mime_type="application/json" if json_uit and not zoeken else "text/plain",
                        tools=[types.Tool(google_search=types.GoogleSearch())] if zoeken else None,
                    )
                    r = client.models.generate_content(model=model, contents=tekst, config=cfg)
                    if not r.text:
                        raise RuntimeError("leeg antwoord")
                    return (r.text, bronnen_gemini(r)) if met_bronnen else r.text
                extra = {"tools": [{"type": "web_search_20250305", "name": "web_search", "max_uses": 5}]} if zoeken else {}
                r = client.messages.create(model=model, max_tokens=min(TOKENS[taak], 8000),
                                           messages=[{"role": "user", "content": tekst}], **extra)
                antw = "".join(b.text for b in r.content if b.type == "text")
                return (antw, bronnen_claude(r)) if met_bronnen else antw
            except Exception as e:
                fout = str(e)
                laatste_fout = fout
                limiet = is_limiet(fout)
                if limiet and is_daglimiet(fout):
                    UITGEPUT.add(model)
                    print(f"  {model}: daglimiet of tegoed op, deze run niet meer gebruiken")
                    break
                if limiet and poging >= 1:
                    print(f"  {model}: limiet bereikt, ander model proberen")
                    break
                wacht = 40 if limiet else 20 * (poging + 1)
                print(f"  {model}: fout ({fout[:110]}), wacht {wacht}s, poging {poging + 1}/3")
                time.sleep(wacht)
    if laatste_fout and is_limiet(laatste_fout):
        raise LimietOp(laatste_fout[:200])
    raise RuntimeError(f"Taalmodel faalt: {(laatste_fout or '')[:200]}")


def parse_json(tekst):
    tekst = re.sub(r"^```(?:json)?|```$", "", tekst.strip(), flags=re.M).strip()
    try:
        return json.loads(tekst)
    except json.JSONDecodeError:  # bijv. tekst rond de JSON na een zoekactie
        m = re.search(r"\{.*\}", tekst, re.S)
        if not m:
            raise
        return json.loads(m.group(0))


# ---------- stap 1: zoeken in het CVDR ----------
def zoek_cvdr(gemeente, trefwoorden=None):
    """Geeft {cvdr_id: {versie, titel, xml_url, gewijzigd}} met alleen de hoogste versie."""
    gevonden = {}
    per_pagina = CFG["max_resultaten_per_zoekvraag"]
    for woord in trefwoorden or CFG["trefwoorden"]:
        for start in range(1, 4 * per_pagina, per_pagina):  # grote gemeenten: max. 4 pagina's
            params = {
                "version": "1.2", "operation": "searchRetrieve", "x-connection": "cvdr",
                "query": f'dcterms.creator="{gemeente}" AND keyword all "{woord}"',
                "startRecord": start, "maximumRecords": per_pagina,
                "x-info-1-accept": "any",
            }
            r = requests.get(SRU, params=params, headers=UA, timeout=60)
            r.raise_for_status()
            dbg = DATA / "debug_sru_voorbeeld.xml"  # bewaar één ruwe respons om veldnamen te checken
            if not dbg.exists():
                dbg.write_text(r.text, encoding="utf-8")
            for rec in re.findall(r"<(?:\w+:)?recordData>(.*?)</(?:\w+:)?recordData>", r.text, re.S):
                m = re.search(r">(CVDR\d+)_(\d+)<", rec)
                url = re.search(r"(https://repository\.officiele-overheidspublicaties\.nl/cvdr/CVDR\d+/\d+/xml/[^<\s\"]+\.xml)", rec, re.I)
                if not (m and url):
                    continue
                eind = re.search(r"uitwerkingtreding[^>]*>(\d{4}-\d{2}-\d{2})<", rec, re.I)
                if eind and eind.group(1) < VANDAAG:
                    continue  # regeling is al vervallen
                cid, versie = m.group(1), int(m.group(2))
                titel = re.search(r"<dcterms:title[^>]*>(.*?)</dcterms:title>", rec, re.S)
                gew = re.search(r"<dcterms:modified[^>]*>(\d{4}-\d{2}-\d{2})", rec)
                if cid not in gevonden or versie > gevonden[cid]["versie"]:
                    gevonden[cid] = {"versie": versie, "xml_url": url.group(1),
                                     "titel": html.unescape(titel.group(1).strip()) if titel else cid,
                                     "gewijzigd": gew.group(1) if gew else None}
            time.sleep(1)  # netjes blijven tegen de overheidsserver
            totaal = re.search(r"numberOfRecords>(\d+)<", r.text)
            if not totaal or int(totaal.group(1)) < start + per_pagina:
                break  # alle resultaten binnen
    return gevonden


def haal_alle_gemeenten():
    """Alle Nederlandse gemeenten via de PDOK Locatieserver. Lukt dat niet, dan de lijst van de vorige run."""
    url = "https://api.pdok.nl/bzk/locatieserver/search/v3_1/free"
    gemeenten = {}
    try:
        for start in range(0, 500, 100):
            r = requests.get(url, params={"q": "*:*", "fq": "type:gemeente", "fl": "gemeentenaam,provincienaam",
                                          "rows": 100, "start": start}, headers=UA, timeout=60)
            r.raise_for_status()
            docs = r.json()["response"]["docs"]
            for d in docs:
                gemeenten[d["gemeentenaam"]] = d.get("provincienaam")
            if len(docs) < 100:
                break
    except Exception as e:
        print(f"Gemeentelijst ophalen bij PDOK mislukt ({e}); vorige lijst gebruiken")
        return lees(ALLE_GEMEENTEN, [])
    lijst = [{"naam": n, "provincie": p} for n, p in sorted(gemeenten.items())]
    if len(lijst) > 300:  # alleen bewaren als de lijst compleet lijkt
        schrijf(ALLE_GEMEENTEN, lijst)
        return lijst
    print(f"Gemeentelijst van PDOK lijkt onvolledig ({len(lijst)}); vorige lijst gebruiken")
    return lees(ALLE_GEMEENTEN, lijst)


def gemeentelijst():
    """Alle gemeenten (via PDOK), of alleen config.json als "alle_gemeenten" uit staat."""
    if not CFG.get("alle_gemeenten", True):
        lijst = CFG["gemeenten"]
    else:
        lijst = haal_alle_gemeenten() or CFG["gemeenten"]
    if ALLEEN:
        lijst = [g for g in lijst if g["naam"].lower() in ALLEEN]
    return lijst


# ---------- stap 5: gemeentesite en partners (en zo nodig AI-zoeken) ----------
WEB_EXTRACTIE = """

LET OP, afwijkend van hierboven: de tekst hieronder komt niet uit het CVDR maar van een of meer WEBPAGINA'S
(gemeentesite of een partner zoals een energieloket) over gemeente {g}. Elke pagina begint met
"=== Pagina: <webadres> ===". Er kunnen meerdere regelingen op staan, of geen enkele.
- Neem alleen regelingen op die geld, korting, een voucher, een gratis isolatieactie of een lening geven voor
  ISOLATIE van BESTAANDE woningen van particulieren, en die gelden voor inwoners van {g}.
  Brede verduurzamingsregelingen en duurzaamheidsleningen tellen mee als isolatie eronder valt.
  Niet: landelijke regelingen (ISDE, Warmtefonds), regelingen van andere gemeenten, alleen monumenten,
  bedrijven of verhuurders, en losse energieadviezen zonder geld of korting.
- Zet bij elke regeling "bron_url" op het webadres van de pagina waar je hem vond.
- Zet "uitvoerder" op wie de regeling uitvoert of de aanvraag behandelt (gemeente of partner).
- Staat er dat het budget op is of de regeling gesloten is: "budgetstatus": "uitgeput".
- Al bekend uit het CVDR of handmatig ingevoerd voor {g}: {bekend}.
  Is een regeling op de pagina dezelfde als een daarvan, zet "zelfde_als" op die naam.
- "bewijs" moet LETTERLIJK uit de pagina's komen.

Geef ALLEEN geldige JSON: {{"regelingen": [ {{ ...per regeling het formaat hierboven, plus "bron_url",
"uitvoerder" en "zelfde_als"... }} ]}}. Geen passende regeling? {{"regelingen": []}}."""


def haal_webpagina(url):
    p = webbronnen.lees_pagina(url)
    if not p:
        raise RuntimeError(f"{url} niet bereikbaar")
    return p["tekst"]


def extraheer_web(gemeente, tekst, bekend):
    vraag = (PROMPT + WEB_EXTRACTIE.format(g=gemeente, bekend="; ".join(bekend) or "geen")
             + "\n\n<webpaginas>\n" + tekst + "\n</webpaginas>")
    antw = parse_json(llm("extractie", vraag, json_uit=True))
    if isinstance(antw, list):
        return antw
    return antw.get("regelingen") or []


ZOEKEN_OP = []   # gevuld zodra zoeken met Google deze run niet meer kan


def ai_zoekfunctie(prompt):
    """AI-zoekactie met Google. Kan dat niet meer (Gemini op) maar werken de reserve-aanbieders nog,
    dan None: de run gaat door zonder AI-zoekacties."""
    if ZOEKEN_OP or client is None:
        return None
    try:
        return llm("zoeken", prompt, zoeken=True, met_bronnen=True)
    except LimietOp as e:
        if not reserve_aanbieders():
            raise
        ZOEKEN_OP.append(str(e))
        print(f"  AI-zoeken kan deze run niet meer ({str(e)[:100]}); verder met site en partners")
        return None


def bewijs_toepassen(rec, tekst):
    """Controle op verzinnen: staat het bewijs van het model echt in de brontekst?"""
    ctrl = webbronnen.controleer_bewijs(rec, tekst)
    rec["bewijs_controle"] = {"datum": VANDAAG, **ctrl}
    if ctrl["niet_gevonden"]:
        rec["betrouwbaarheid"] = "laag" if rec.get("bron") == "web" else "middel"
        rec["opmerkingen"] = (f"Bewijs niet teruggevonden in de brontekst voor: {', '.join(ctrl['niet_gevonden'])}. "
                              "Controleer deze velden. " + (rec.get("opmerkingen") or "")).strip()
    return rec


def web_records(g, ext_lijst, tekst, verslag, oud_web, bestaand):
    """Zet de uitkomst van de web-extractie om naar records.
    bestaand: CVDR/handmatige records van deze gemeente (dezelfde regeling niet dubbel opnemen).
    Geeft (nieuwe records, [(bestaand record, webadres)] voor 'ook vermeld op', logregels)."""
    naam = g["naam"]
    recs, ook, log = [], [], []
    bronnen = {p["url"]: p["bron"] for p in verslag["paginas"]}
    for ext in ext_lijst:
        if not isinstance(ext, dict) or not ext.get("relevant", True):
            continue
        titel = (ext.get("naam") or "").strip()
        if not titel:
            continue
        url = ext.get("bron_url") if ext.get("bron_url") in bronnen else (verslag["paginas"][0]["url"] if verslag["paginas"] else None)
        zelfde = next((r for r in bestaand if (ext.get("zelfde_als") and webbronnen.lijkt_op(ext["zelfde_als"], r["naam"], naam))
                       or webbronnen.lijkt_op(titel, r["naam"], naam)), None)
        if zelfde:
            ook.append((zelfde, url))
            continue
        vorige = next((r for r in oud_web if webbronnen.lijkt_op(titel, r["naam"], naam)), None)
        rec = {
            "id": vorige["id"] if vorige else f"{slug(naam)}-web-{slug(titel)[:50]}", "cvdr_id": None, "bron": "web",
            "bron_type": bronnen.get(url, "internet"), "handmatig": False, "gecontroleerd": False,
            "gemeente": naam, "provincie": g.get("provincie"), "peildatum": VANDAAG,
        }
        for k in ("relevant", "zelfde_als"):
            ext.pop(k, None)
        rec.update({k: v for k, v in ext.items() if v is not None})
        rec["naam"], rec["bron_url"] = titel, url
        rec = normaliseer(rec)
        rec["status"] = status(rec)
        if rec.get("budgetstatus") == "uitgeput" and rec["status"] != "gesloten":
            rec["status"] = "gesloten"
        rec["betrouwbaarheid"] = "middel"
        rec["opmerkingen"] = (f"Gevonden op internet ({rec['bron_type']}), niet in het CVDR. Controleer de bronpagina. "
                              + (rec.get("opmerkingen") or "")).strip()
        rec = bewijs_toepassen(rec, tekst)
        if any(r["id"] == rec["id"] for r in recs):
            continue
        recs.append(rec)
        diff = verschillen(vorige, rec)
        if not vorige or diff:
            kop = f"**{naam} – {rec['naam']}** ({'nieuw' if not vorige else 'gewijzigd'} via {rec['bron_type']}, {url})"
            log.append(kop + "".join(f"\n  - {v}: {json.dumps(a, ensure_ascii=False)} → {json.dumps(b, ensure_ascii=False)}"
                                     for v, a, b in diff if vorige))
        if rec["bewijs_controle"]["niet_gevonden"]:
            log.append(f"**{naam} – {rec['naam']}**: bewijs niet teruggevonden voor "
                       f"{', '.join(rec['bewijs_controle']['niet_gevonden'])} (controleren)")
    return recs, ook, log


# ---------- stap 2: tekst ophalen ----------
def haal_tekst(xml_url, pagina_url=None):
    """Tekst van de regeling uit de XML. Geeft die een fout (bijv. 404), dan de gewone CVDR-pagina."""
    try:
        r = requests.get(xml_url, headers=UA, timeout=60)
        r.raise_for_status()
    except Exception as e:
        if not pagina_url:
            raise
        print(f"  XML niet bereikbaar ({e}); CVDR-pagina proberen")
        return haal_webpagina(pagina_url)[:150_000]
    t = re.sub(r"<[^>]+>", " ", r.text)
    t = re.sub(r"\s+", " ", html.unescape(t)).strip()
    if "<illustratie" in r.text or ".jpg" in r.text or ".png" in r.text:
        t += "\n[LET OP: deze regeling bevat afbeeldingen (bijv. een bijlage) die niet als tekst zijn meegenomen.]"
    return t[:150_000]


# ---------- stap 3: trechter (eerst gratis, dan pas Gemini) ----------
# Titels die nooit een isolatieregeling voor woningen zijn
TITEL_NIET = re.compile(
    r"algemene plaatselijke|\bapv\b|bouwverordening|leges|omgevingsplan|bestemmingsplan|mandaat|delegatie|"
    r"volmacht|welstand|monument|erfgoed|afval|precario|"
    r"tarie|belasting|huisvesting|parkeer|evenement|sport|cultuur|onderwijs|jeugd|wmo|bijstand|participatie|"
    r"inspraak|klacht|archief|begroting|reglement van orde|horeca|kinderopvang|verkeer|riool|"
    r"dienstverlening|aanwijzingsbesluit|restauratie|bedrijven|maatschappelijk vastgoed|subsidieplafond", re.I)
# Titels die meestal over iets anders gaan, behalve als isolatie of verduurzaming in de titel staat
TITEL_MEESTAL_NIET = re.compile(r"warmtepomp|zonne|warmtenet|groene? da|water", re.I)
TITEL_ISOLATIE = re.compile(r"isol|verduurzam|energiebespar|energiezuinig|duurzaam(heids)?lening|duurzame? woning|"
                            r"aardgasvrij|energietransitie|woningverbetering", re.I)
TITEL_WEL = re.compile(r"subsidie|regeling|regels|verordening|lening|voucher|waardebon|tegoed|bijdrage|stimulering|"
                       r"isol|glas|fonds|verduurzam|energie|duurzaam", re.I)
# Eerder door het (strengere) AI-filter afgewezen titels die met het ruimere filter opnieuw bekeken worden
TITEL_HERBEKIJKEN = re.compile(r"lening|verduurzam|duurzaam|aardgasvrij|energie|fonds|stimulering|woning", re.I)


def titel_valt_af(titel, gemeente=""):
    """Stap 3a: alleen op de titel, zonder iets te downloaden.
    De gemeentenaam telt niet mee (anders valt bijv. alles van 'Waterland' af op 'water')."""
    if gemeente:
        titel = re.sub(re.escape(gemeente), " ", titel, flags=re.I)
    if TITEL_NIET.search(titel) or not TITEL_WEL.search(titel):
        return True
    return bool(TITEL_MEESTAL_NIET.search(titel)) and not TITEL_ISOLATIE.search(titel)


def is_vervallen(cid):
    """Stap 3b: de CVDR-pagina toont 'Geldend van ... t/m <datum>'. Ligt die datum in het verleden: overslaan."""
    try:
        r = requests.get(f"https://lokaleregelgeving.overheid.nl/{cid}", headers=UA, timeout=60)
        t = re.sub(r"<[^>]+>", " ", r.text)
        m = re.search(r"Geldend van\s+\d{2}-\d{2}-\d{4}\s+t/m\s+(heden|\d{2}-\d{2}-\d{4})", t)
        time.sleep(0.5)
        if not m or m.group(1) == "heden":
            return False
        d, mnd, j = m.group(1).split("-")
        return f"{j}-{mnd}-{d}" < VANDAAG
    except Exception:
        return False  # bij twijfel niet overslaan


def tekst_valt_af(tekst):
    """Stap 3c: moet echt over isolatie van woningen en geld gaan."""
    laag = tekst.lower()
    return not ("isol" in laag and ("woning" in laag or "eigenaar" in laag)
                and re.search(r"subsidie|lening|tegoed|voucher|waardebon|bijdrage", laag))


def is_relevant(titel, tekst):
    """Stap 3d: pas nu het taalmodel (Flash-Lite) vragen."""
    vraag = ("Beantwoord met alleen JA of NEE.\n"
             "JA als particuliere eigenaren met deze gemeentelijke regeling geld, korting, een voucher of een lening "
             "kunnen krijgen voor ISOLATIE van hun BESTAANDE woning: dak, zolder, gevel, spouwmuur, vloer/bodem of "
             "isolerend glas (HR++/triple). Ook JA als isolatie één van de maatregelen is in een bredere "
             "verduurzamingsregeling, aardgasvrij-regeling of duurzaamheidslening.\n"
             "NEE als isolatie er niet onder valt (bijv. alleen zonnepanelen, warmtepomp, groen dak, afkoppelen "
             "regenwater), als het alleen om monumenten, bedrijven, verenigingen of verhuurders gaat, of als het "
             "geen regeling met geld is (bijv. een besluit over een subsidieplafond of mandaat).\n\n"
             f"Titel: {titel}\n\nBegin van de tekst:\n{tekst[:6000]}")
    return llm("filter", vraag).strip().upper().startswith("JA")


# ---------- stap 4: voorwaarden extraheren ----------
def extraheer(tekst):
    antw = llm("extractie", PROMPT + "\n\n<regelingstekst>\n" + tekst + "\n</regelingstekst>", json_uit=True)
    return parse_json(antw)


TOETS = ["eigenaar_bewoner", "woz_max", "isolatiestaat", "bouwjaar_max", "woonoppervlak_max", "inkomen", "vve"]


def betrouwbaarheid(rec):
    if re.search(r"tegenstrijd|spreekt.*tegen", rec.get("opmerkingen") or "", re.I):
        return "middel"
    if not rec.get("bedrag") or not rec.get("looptijd_eind"):
        return "middel"
    return "hoog"


def status(rec):
    eind = rec.get("looptijd_eind")
    if eind and eind < VANDAAG:
        return "gesloten"
    return "open" if eind else "onbekend"


# ---------- wijzigingen ----------
VELDEN_LOG = ["bedrag", "looptijd_eind", "aanvragen", "inkomensgrens", "stapelbaar_isde", "status"]


def verschillen(oud, nieuw):
    uit = []
    for v in VELDEN_LOG:
        if (oud or {}).get(v) != nieuw.get(v):
            uit.append((v, (oud or {}).get(v), nieuw.get(v)))
    oc, nc = (oud or {}).get("criteria") or {}, nieuw.get("criteria") or {}
    for v in sorted(set(oc) | set(nc)):
        if oc.get(v) != nc.get(v):
            uit.append((f"criteria.{v}", oc.get(v), nc.get(v)))
    return uit


# ---------- hoofdprogramma ----------
def main():
    if "--alleen-overzicht" in sys.argv:  # alleen de sheet opnieuw maken uit de bestaande bestanden
        schrijf_overzicht(lees(OUT, []), lees(ZOEKSTATUS, {}).get("gemeenten", []), lees(DEKKING, {}))
        export_tios.schrijf(ROOT)
        return
    state = {} if FORCEER else lees(STATE, {})
    web_state = lees(WEB_STATE, {})
    oud = {r["id"]: normaliseer(r) for r in lees(OUT, [])}
    oud_per_cvdr = {r["cvdr_id"]: r for r in oud.values() if r.get("cvdr_id")}
    dekking = lees(DEKKING, {})
    zoekstatus = {g["gemeente"]: g for g in lees(ZOEKSTATUS, {}).get("gemeenten", [])}
    nieuw, log = [], []
    limiet_op, uitgesteld, web_gedaan = False, 0, 0
    web_interval = CFG.get("web_zoeken_elke_dagen", 30)
    web_max = CFG.get("max_web_zoekacties_per_run", 25)
    gemeenten = gemeentelijst()
    namen = {g["naam"] for g in gemeenten}
    sites = lees(GEMEENTE_SITES, {})
    roo = {}
    partners = webbronnen.Partners(CFG.get("partners"))

    def site_van(naam):
        """Webadres van de gemeentesite, bewaard in data/gemeente_sites.json (elke 90 dagen opnieuw opgezocht)."""
        bekend = sites.get(naam) or {}
        if bekend.get("datum") and (dt.date.fromisoformat(VANDAAG) - dt.date.fromisoformat(bekend["datum"])).days < 90:
            return bekend.get("url")
        if not roo:
            roo.update(webbronnen.haal_roo() or {"_leeg": ""})
        sites[naam] = {"url": webbronnen.vind_site(naam, roo, CFG.get("gemeente_sites")), "datum": VANDAAG}
        return sites[naam]["url"]

    print(f"{len(gemeenten)} gemeenten")
    # Is de vorige run halverwege gestopt? Dan beginnen bij de gemeente waar hij bleef.
    volgende = None if ALLEEN else lees(VOORTGANG, {}).get("volgende")
    if volgende in namen:
        i = [g["naam"] for g in gemeenten].index(volgende)
        gemeenten = gemeenten[i:] + gemeenten[:i]
        print(f"Verder bij {volgende} (daar stopte de vorige run)")
    # PR-tekst van een openstaande Subsidie-update, voordat deze run hem overschrijft
    vorige_tekst = (LOG_LAST.read_text(encoding="utf-8")
                    if os.environ.get("VERDER_OP_OPEN_UPDATE") == "1" and LOG_LAST.exists() else "")
    gedaan = set()
    stand = {"volgende": None, "reden": None}
    resultaten = []   # (gemeente, symbool, tekst) voor de samenvatting

    def resultaat(naam, sym, tekst):
        print("::endgroup::")
        print(f"{sym} {naam}: {tekst}")
        resultaten.append((naam, sym, tekst))

    def pr_tekst():
        tekst = f"# Subsidie-update {VANDAAG}\n\n" + ("\n\n".join(log) if log else "Geen wijzigingen.") + "\n"
        wacht = sum(len(z["wacht_op_uitlezen"]) for z in zoekstatus.values())
        if wacht:
            tekst += (f"\n**Let op:** nog {wacht} regeling(en) in het CVDR wachten op uitlezen; "
                      "die volgen bij de volgende run(s).\n")
        if stand["volgende"]:
            tekst += (f"\n**Let op:** deze run is gestopt ({stand['reden']}) na {len(gedaan)} van {len(namen)} gemeenten. "
                      f"De volgende run gaat verder bij {stand['volgende']}.\n")
        # Gaat deze run verder op een openstaande Subsidie-update? Dan blijven de wijzigingen van de
        # eerdere runs in de PR-tekst staan (zonder hun oude "Let op"-regels).
        if vorige_tekst:
            eerder = re.sub(r"\n\*\*Let op:\*\*[^\n]*\n?", "\n", vorige_tekst)
            eerder = re.sub(r"^# Subsidie-update", "### Subsidie-update", eerder, flags=re.M)
            eerder = eerder.replace("## Eerdere runs in deze update (nog niet gemerged)", "")
            eerder = re.sub(r"\n{3,}", "\n\n", eerder).strip()
            tekst += "\n## Eerdere runs in deze update (nog niet gemerged)\n\n" + eerder + "\n"
            if len(tekst) > 60000:   # GitHub staat max. 65.536 tekens toe in een PR-tekst
                tekst = tekst[:60000] + "\n\n… (ingekort; alles staat in data/wijzigingen.md)\n"
        return tekst

    def opslaan(definitief=False):
        """Schrijft alles weg wat tot nu toe gedaan is. Gemeenten die nog niet aan bod kwamen
        houden hun vorige uitkomst. Kan op elk moment, dus ook tussentijds en bij een fout."""
        in_nieuw = {r["id"] for r in nieuw}
        lijst = nieuw + [r for r in oud.values() if r.get("handmatig")]
        lijst += [r for r in oud.values() if not r.get("handmatig") and r["gemeente"] not in gedaan
                  and r["id"] not in in_nieuw]
        lijst = [normaliseer(r) for r in {r["id"]: r for r in lijst}.values()]
        lijst.sort(key=lambda r: (r["gemeente"], r["naam"]))
        schrijf(OUT, lijst)
        schrijf(STATE, state)
        schrijf(WEB_STATE, web_state)
        schrijf(GEMEENTE_SITES, sites)
        schrijf(DEKKING, dekking)
        schrijf(ZOEKSTATUS, {"peildatum": VANDAAG,
                             "gemeenten": sorted(zoekstatus.values(), key=lambda x: x["gemeente"])})
        if not ALLEEN:
            schrijf(VOORTGANG, {"volgende": stand["volgende"], "datum": VANDAAG})
        tekst = pr_tekst()
        LOG_LAST.write_text(tekst, encoding="utf-8")
        if definitief:
            with LOG_ALL.open("a", encoding="utf-8") as f:
                f.write("\n" + tekst)
        return lijst

    def stop_signaal(signum, frame):
        raise KeyboardInterrupt(f"signaal {signum}")
    signal.signal(signal.SIGTERM, stop_signaal)
    laatst_bewaard = time.monotonic()

    try:
      for i_g, g in enumerate(gemeenten):
        if time.monotonic() - START > MAX_RUN_SEC:
            stand.update(volgende=g["naam"], reden="tijd van de run op")
            print(f"Tijd van de run op: stoppen. Volgende run begint bij {g['naam']}.")
            break
        naam = g["naam"]
        print(f"::group::{naam} (details)")
        try:
            treffers = zoek_cvdr(naam)
            if not treffers and "(" in naam:  # PDOK zegt "Hengelo (O)", het CVDR "Hengelo"
                treffers = zoek_cvdr(re.sub(r"\s*\(.*?\)", "", naam).strip())
        except Exception as e:
            # CVDR even niet bereikbaar: vorige CVDR-uitkomst houden en wel op internet verder zoeken
            print(f"  CVDR-zoekvraag mislukt ({str(e)[:150]}); vorige uitkomst behouden")
            cvdr_fout = True
            treffers = {}
            nieuw += [r for r in oud.values() if r["gemeente"] == naam and r.get("cvdr_id") and not r.get("handmatig")]
        else:
            cvdr_fout = False
        print(f"  {len(treffers)} treffers in CVDR")
        relevant, wachtrij = 0, []
        vorige_status = zoekstatus.get(naam, {})
        if cvdr_fout:   # wachtrij van de vorige keer niet kwijtraken
            wachtrij = list(vorige_status.get("wacht_op_uitlezen") or [])
        tel = {"titel": 0, "vervallen": 0, "geen isolatie": 0, "gemini-filter": 0, "ongewijzigd": 0, "uitgelezen": 0}
        for cid, meta in treffers.items():
            st = state.get(cid, {})
            vorige = oud_per_cvdr.get(cid)
            zelfde_versie = not FORCEER and st.get("versie") == meta["versie"]

            # al eerder beoordeeld en niets veranderd: niets downloaden
            if zelfde_versie and st.get("skip"):
                # ruimere filters (FILTER_VERSIE): eerder afgewezen titels die er nu wel door kunnen, opnieuw bekijken
                herbekijk = st.get("fv", 1) < FILTER_VERSIE and (
                    (st["skip"] == "titel" and not titel_valt_af(meta["titel"], naam))
                    or (st["skip"] == "gemini-filter" and TITEL_HERBEKIJKEN.search(meta["titel"])))
                if not herbekijk:
                    tel[st["skip"]] = tel.get(st["skip"], 0) + 1
                    continue
            if zelfde_versie and vorige:
                nieuw.append({**vorige, "peildatum": VANDAAG, "status": status(vorige)})
                relevant += 1
                tel["ongewijzigd"] += 1
                continue

            def overslaan(reden):
                state[cid] = {"versie": meta["versie"], "skip": reden, "fv": FILTER_VERSIE}
                tel[reden] += 1

            def later():  # limiet bereikt: bewaren voor de volgende run
                nonlocal uitgesteld
                uitgesteld += 1
                wachtrij.append({"cvdr_id": cid, "titel": meta["titel"], "gewijzigd": meta["gewijzigd"],
                                 "bron_url": f"https://lokaleregelgeving.overheid.nl/{cid}/{meta['versie']}"})
                if vorige:
                    nieuw.append(vorige)

            if titel_valt_af(meta["titel"], naam):
                overslaan("titel"); continue
            if is_vervallen(cid):
                overslaan("vervallen"); continue
            try:
                tekst = haal_tekst(meta["xml_url"], f"https://lokaleregelgeving.overheid.nl/{cid}/{meta['versie']}")
            except Exception as e:
                print(f"  {cid}: tekst ophalen mislukt ({e})")
                later(); continue
            if tekst_valt_af(tekst):
                overslaan("geen isolatie"); continue
            tekst_hash = hashlib.sha256(tekst.encode()).hexdigest()
            if vorige and not FORCEER and st.get("hash") == tekst_hash:
                # alleen het versienummer is veranderd, de tekst niet: vorige uitkomst houden
                state[cid] = {"versie": meta["versie"], "hash": tekst_hash}
                nieuw.append({**vorige, "versie": meta["versie"], "peildatum": VANDAAG, "status": status(vorige),
                              "bron_url": f"https://lokaleregelgeving.overheid.nl/{cid}/{meta['versie']}"})
                relevant += 1
                tel["ongewijzigd"] += 1
                continue
            try:
                print(f"  {cid}: filter ({meta['titel'][:60]})")
                if not is_relevant(meta["titel"], tekst):
                    overslaan("gemini-filter"); continue
                print(f"  {cid} v{meta['versie']}: uitlezen")
                ext = extraheer(tekst)
            except ZonderAI:
                later(); continue
            except LimietOp as e:
                print(f"  AI-BUDGET OF LIMIET OP ({e}). Run stopt; de volgende gaat hier verder.")
                limiet_op = limiet_reden(e)
                break
            except Exception as e:
                print(f"  {cid}: mislukt ({e})")
                later(); continue
            if not ext.get("relevant", True):
                overslaan("gemini-filter"); continue
            ext.pop("relevant", None)
            rec = {
                "id": vorige["id"] if vorige else f"{slug(naam)}-{cid.lower()}", "cvdr_id": cid, "versie": meta["versie"],
                "bron": "cvdr", "handmatig": False, "gecontroleerd": False,
                "gemeente": naam, "provincie": g.get("provincie"),
                **ext,
                "bron_url": f"https://lokaleregelgeving.overheid.nl/{cid}/{meta['versie']}",
                "datum_regelingstekst": meta.get("gewijzigd"), "peildatum": VANDAAG,
            }
            rec = normaliseer(rec)
            rec["naam"] = rec.get("naam") or meta["titel"]
            rec["status"] = status(rec)
            rec["betrouwbaarheid"] = betrouwbaarheid(rec)
            rec = bewijs_toepassen(rec, tekst)
            state[cid] = {"versie": meta["versie"], "hash": tekst_hash}
            relevant += 1
            tel["uitgelezen"] += 1
            nieuw.append(rec)
            diff = verschillen(vorige, rec)
            kop = f"**{naam} – {rec['naam']}** ({'nieuw' if not vorige else 'gewijzigd'}, {rec['bron_url']})"
            log.append(kop + "".join(f"\n  - {v}: {json.dumps(a, ensure_ascii=False)} → {json.dumps(b, ensure_ascii=False)}" for v, a, b in diff))
            if rec["bewijs_controle"]["niet_gevonden"]:
                log.append(f"**{naam} – {rec['naam']}**: bewijs niet teruggevonden voor "
                           f"{', '.join(rec['bewijs_controle']['niet_gevonden'])} (controleren)")

        print("  trechter: " + ", ".join(f"{k} {v}" for k, v in tel.items() if v))
        if limiet_op:  # budget/limiet op: deze gemeente is niet af, de volgende run begint hier
            stand.update(volgende=naam, reden=limiet_op)
            resultaat(naam, "🛑", f"{limiet_op}; de volgende run gaat hier verder")
            break

        # regelingen die niet meer gevonden worden: niet weggooien, wel markeren
        gezien = {r["cvdr_id"] for r in nieuw if r.get("cvdr_id")}
        for r in ([] if cvdr_fout else oud.values()):
            if r["gemeente"] == naam and r.get("cvdr_id") and r["cvdr_id"] not in gezien and not r.get("handmatig"):
                r = {**r, "status": "onbekend", "gecontroleerd": False,
                     "opmerkingen": f"Niet meer gevonden in CVDR op {VANDAAG}: mogelijk ingetrokken of vervangen. " + (r.get("opmerkingen") or "")}
                nieuw.append(r)
                log.append(f"**{naam} – {r['naam']}**: niet meer gevonden in CVDR")

        # stap 5: voor elke gemeente ook de gemeentesite en partners (en zo nodig AI-zoeken)
        oud_web = [r for r in oud.values() if r["gemeente"] == naam and r.get("bron") == "web"]
        ws = web_state.get(naam, {})
        web_te_oud = (FORCEER or ws.get("versie") != WEB_VERSIE or not ws.get("datum")
                      or (dt.date.fromisoformat(VANDAAG) - dt.date.fromisoformat(ws["datum"])).days >= web_interval)
        web_uitkomst = None   # wat het zoeken op internet deze run opleverde
        web_paginas = 0
        if web_te_oud:
            bestaand = ([r for r in nieuw if r["gemeente"] == naam and r.get("bron") != "web"]
                        + [r for r in oud.values() if r["gemeente"] == naam and r.get("handmatig")])
            heeft_open = any(r.get("status") != "gesloten" for r in bestaand)
            mag_ai_zoeken = not ZONDER_AI and web_gedaan < web_max
            try:
                print("  zoeken op de gemeentesite en bij partners")
                gev = webbronnen.verzamel(naam, g.get("provincie"), site_van(naam), partners,
                                          zoek_functie=ai_zoekfunctie if mag_ai_zoeken else None, parse_json=parse_json,
                                          ai_altijd=CFG.get("ai_zoeken_altijd", False) or not heeft_open)
                vs = gev["verslag"]
                web_paginas = len(vs["paginas"])
                web_gedaan += vs["ai_gezocht"]
                print(f"  site: {vs['site'] or 'onbekend'} ({vs['site_bron'] or '-'}, {vs['site_urls']} adressen), "
                      f"{vs['kandidaten']} kandidaten, {len(vs['paginas'])} relevante pagina('s)"
                      + (", AI-zoekactie gedaan" if vs["ai_gezocht"] else ""))
                for pg in vs["paginas"]:
                    print(f"    {pg['bron']}: {pg['url']}")
                h = webbronnen.tekst_hash(gev["paginas"]) if gev["paginas"] else None
                nieuw_ws = {"versie": WEB_VERSIE, "datum": VANDAAG, "site": vs["site"], "site_bron": vs["site_bron"],
                            "site_urls": vs["site_urls"], "kandidaten": vs["kandidaten"], "partners": vs["partners"],
                            "ai_gezocht": VANDAAG if vs["ai_gezocht"] else ws.get("ai_gezocht"),
                            "paginas": vs["paginas"], "hash": h, "ook": ws.get("ook", [])}
                if not gev["paginas"]:
                    web_recs = oud_web   # niets (meer) gevonden of niet bereikbaar: vorige uitkomst blijft staan
                elif ZONDER_AI:
                    web_recs = oud_web
                elif h == ws.get("hash") and not FORCEER:
                    print("  pagina's ongewijzigd: vorige uitkomst blijft")
                    web_recs = [{**r, "peildatum": VANDAAG} for r in oud_web]
                else:
                    tekst = webbronnen.samengevoegd(gev["paginas"])
                    print(f"  uitlezen ({len(tekst)} tekens uit {len(gev['paginas'])} pagina('s))")
                    ext = extraheer_web(naam, tekst, [r["naam"] for r in bestaand])
                    web_recs, ook, wlog = web_records(g, ext, tekst, vs, oud_web, bestaand)
                    log += wlog
                    nieuw_ws["ook"] = sorted({(r["id"], u) for r, u in ook if u})
                    ids = {r["id"] for r in web_recs}
                    for r in oud_web:   # niet meer gevonden: niet weggooien, wel markeren
                        if r["id"] in ids:
                            continue
                        if not (r.get("opmerkingen") or "").startswith("Niet meer gevonden op internet"):
                            log.append(f"**{naam} – {r['naam']}**: niet meer gevonden op internet")
                            r = {**r, "status": "onbekend", "opmerkingen": f"Niet meer gevonden op internet op {VANDAAG}. "
                                 + (r.get("opmerkingen") or "")}
                        web_recs.append(r)
                    print(f"  internet: {len(web_recs)} regeling(en), {len(ook)} al bekend uit het CVDR")
                if not ZONDER_AI:   # zonder AI niets bewaren: de volgende echte run moet deze gemeente nog uitlezen
                    web_state[naam] = nieuw_ws
                nieuw += web_recs
                web_uitkomst = "gezocht"
            except LimietOp as e:
                print(f"  AI-BUDGET OF LIMIET OP bij zoeken op internet ({e}). Run stopt.")
                limiet_op = limiet_reden(e)
                nieuw += oud_web
                stand.update(volgende=naam, reden=limiet_op)
                resultaat(naam, "🛑", f"{limiet_op} bij zoeken op internet; de volgende run gaat hier verder")
                break
            except Exception as e:
                print(f"  zoeken op internet mislukt ({type(e).__name__}: {e})")
                nieuw += oud_web
                web_uitkomst = "mislukt"
        else:
            nieuw += oud_web  # vorige internetuitkomst blijft staan tot de volgende zoekronde
        # "ook vermeld op": webpagina's die een CVDR-regeling van deze gemeente noemen
        for rid, url in web_state.get(naam, {}).get("ook", []):
            for r in nieuw:
                if r["id"] == rid and url not in (r.get("extra_bronnen") or []):
                    r["extra_bronnen"] = (r.get("extra_bronnen") or []) + [url]

        webtreffers = sum(1 for r in nieuw if r["gemeente"] == naam and r.get("bron") == "web")
        w = web_state.get(naam, {})
        gecheckt = ["CVDR"]
        if w.get("versie") == WEB_VERSIE:
            gecheckt += (["gemeentesite"] if w.get("site_urls") else [])
            gecheckt += ["partners"] + (["AI-zoekactie"] if w.get("ai_gezocht") else [])
        dekking[naam] = {"datum": VANDAAG, "bronnen_gecheckt": gecheckt,
                         "treffers": len(treffers), "relevante_regelingen": relevant + webtreffers,
                         "wacht_op_uitlezen": len(wachtrij)}
        zoekstatus[naam] = {"gemeente": naam, "provincie": g.get("provincie"),
                            "cvdr_gecheckt": vorige_status.get("cvdr_gecheckt") if cvdr_fout else VANDAAG,
                            "wacht_op_uitlezen": wachtrij,
                            "internet_gezocht": w.get("datum") if w.get("versie") == WEB_VERSIE else None,
                            "gemeentesite": w.get("site"), "gemeentesite_doorzocht": bool(w.get("site_urls")),
                            "partners": w.get("partners") or [], "ai_gezocht": w.get("ai_gezocht"),
                            "paginas_gevonden": len(w.get("paginas") or [])}
        regs = [r for r in nieuw if r["gemeente"] == naam and r.get("status") != "gesloten"]
        zonder = [r for r in regs if not heeft_voorwaarden(r)]
        web_regs = [r for r in regs if r.get("bron") == "web"]
        delen = []
        if tel["uitgelezen"]:
            delen.append(f"{tel['uitgelezen']} nieuw/gewijzigd uitgelezen")
        if tel["ongewijzigd"]:
            delen.append(f"{tel['ongewijzigd']} ongewijzigd")
        if web_regs:
            delen.append(f"{len(web_regs)} via internet (controleren)")
        uitleg = f" ({', '.join(delen)})" if delen else ""
        if cvdr_fout:
            resultaat(naam, "❌", "CVDR niet bereikbaar, vorige uitkomst blijft staan"
                      + ("; internet wel doorzocht" if web_uitkomst == "gezocht" else ""))
        elif web_uitkomst == "mislukt":
            resultaat(naam, "❌", "zoeken op internet mislukt" + (f"; wel {len(regs)} open regeling(en)" if regs else ""))
        elif wachtrij and ZONDER_AI:
            resultaat(naam, "🔎", f"{len(wachtrij)} CVDR-regeling(en) en {web_paginas} "
                                  f"internetpagina('s) gevonden; uitlezen volgt in een run met AI")
        elif wachtrij:
            resultaat(naam, "⚠️", f"{len(wachtrij)} regeling(en) niet uitgelezen (fout bij ophalen of uitlezen), "
                                  f"volgt bij de volgende run" + (f"; wel gelukt: {len(regs)}" if regs else ""))
        elif regs and zonder:
            resultaat(naam, "⚠️", f"{len(regs)} open regeling(en), waarvan {len(zonder)} zonder voorwaarden{uitleg}")
        elif regs:
            resultaat(naam, "✅", f"voorwaarden opgehaald voor {len(regs)} open regeling(en){uitleg}")
        elif any(r["gemeente"] == naam for r in nieuw):
            resultaat(naam, "➖", "alleen gesloten regelingen")
        elif web_uitkomst == "gezocht":
            resultaat(naam, "➖", "geen regeling in CVDR, en op gemeentesite, bij partners en via AI niets gevonden")
        else:
            resultaat(naam, "➖", "geen isolatieregeling gevonden")
        gedaan.add(naam)
        stand["volgende"] = gemeenten[i_g + 1]["naam"] if i_g + 1 < len(gemeenten) else None
        stand["reden"] = "onderbroken"
        if time.monotonic() - laatst_bewaard > BEWAAR_ELKE_SEC:
            opslaan()
            laatst_bewaard = time.monotonic()
            print(f"  (tussentijds opgeslagen: {len(gedaan)} gemeenten gedaan)")
      else:
        stand.update(volgende=None, reden=None)   # alle gemeenten gedaan: volgende keer weer vooraan
    except BaseException as e:
        # Crash, annuleren of signaal: bewaren wat er is, zodat de volgende run verder kan.
        print("::endgroup::")
        print(f"Run onderbroken ({type(e).__name__}: {e}). Opslaan wat er is.")
        if stand["volgende"] is None or stand["reden"] == "onderbroken":
            huidig = [g["naam"] for g in gemeenten if g["naam"] not in gedaan]
            stand.update(volgende=huidig[0] if huidig else None, reden=f"onderbroken: {type(e).__name__}")
        lijst = opslaan(definitief=True)
        try:
            schrijf_overzicht(lijst, lees(ZOEKSTATUS, {}).get("gemeenten", []), dekking)
            export_tios.schrijf(ROOT)
        except Exception as e2:
            print(f"Overzicht/export na onderbreking mislukt ({e2})")
        schrijf_samenvatting(resultaten, stand, len(namen))
        raise

    nieuw = opslaan(definitief=True)
    schrijf_samenvatting(resultaten, stand, len(namen))
    wacht = sum(len(z["wacht_op_uitlezen"]) for z in zoekstatus.values())
    print(f"Klaar: {len(nieuw)} regelingen, {len(log)} wijzigingen, {web_gedaan} zoekacties op internet, {wacht} wachten nog")
    schrijf_overzicht(nieuw, lees(ZOEKSTATUS, {}).get("gemeenten", []), dekking)
    export_tios.schrijf(ROOT)

    if "--vergelijk" in sys.argv:
        vergelijk(pathlib.Path(sys.argv[sys.argv.index("--vergelijk") + 1]), nieuw)


# ---------- overzicht in Excel ----------
def _eur(x):
    try:
        return f"€ {float(x):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return str(x)


def _tekst(v):
    if v is None or v == "" or v == [] or v == {}:
        return ""
    if isinstance(v, bool):
        return "ja" if v else "nee"
    if isinstance(v, list):
        return ", ".join(_tekst(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_tekst(x)}" for k, x in v.items() if _tekst(x))
    return str(v)


def _datum(s):
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s or "")
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else _tekst(s)


def _bedrag(b):
    """Bedrag als tekst; het model geeft soms {max, eenheid, type} in plaats van een zin."""
    if isinstance(b, dict):
        delen = [f"max. {_eur(b['max'])}" if b.get("max") is not None else "",
                 f"min. {_eur(b['min'])}" if b.get("min") is not None else "",
                 f"{b['percentage']}%" if b.get("percentage") is not None else "",
                 _tekst({k: v for k, v in b.items() if k not in ("max", "min", "percentage", "eenheid", "type")})]
        return ", ".join(x for x in delen if x) or _tekst(b)
    return _tekst(b)


def rij_voor(r):
    """Eén regeling als leesbare rij (voorwaarden in gewone taal)."""
    c = r.get("criteria") or {}
    woz = ""
    if c.get("woz_max") is not None:
        woz = ("≤ " if c.get("woz_regel") == "lte" else "< ") + _eur(c["woz_max"])
        woz += f" (peildatum {c['woz_peildatum']})" if c.get("woz_peildatum") else ""
        woz += f". Uitzondering: {c['woz_uitzondering']}" if c.get("woz_uitzondering") else ""
    iso = c.get("isolatiestaat") or {}
    bouwjaar = " ".join(x for x in [f"vanaf {c['bouwjaar_min']}" if c.get("bouwjaar_min") else "",
                                    f"t/m {c['bouwjaar_max']}" if c.get("bouwjaar_max") else "",
                                    _tekst(c.get("bouwjaar_opmerking"))] if x)
    if r.get("handmatig"):
        bron = "handmatig"
    elif r.get("bron") == "web":
        bron = r.get("bron_type") or "internet (AI)"
    else:
        bron = {"wacht": "CVDR (wacht)"}.get(r.get("bron"), "CVDR")
    bc = r.get("bewijs_controle") or {}
    bewijs = ("" if not bc else "niet gevonden: " + ", ".join(bc["niet_gevonden"]) if bc.get("niet_gevonden")
              else f"ok ({bc.get('gecontroleerd', 0)} velden)")
    return {
        "Gemeente": r.get("gemeente"), "Provincie": r.get("provincie"), "Regeling": r.get("naam"),
        "Bron": bron, "Status": r.get("status"), "Gecontroleerd": "" if r.get("bron") == "wacht" else ("ja" if r.get("gecontroleerd") else "nee"),
        "Betrouwbaarheid": r.get("betrouwbaarheid"), "Bedrag": _bedrag(r.get("bedrag")),
        "Aanvragen t/m": _datum(r.get("looptijd_eind")), "Eigenaar-bewoner": _tekst(c.get("eigenaar_bewoner")),
        "WOZ": woz, "Isolatiestaat / energielabel": _tekst(iso.get("omschrijving")) if isinstance(iso, dict) else _tekst(iso),
        "Bouwjaar": bouwjaar, "Max. woonoppervlak": _tekst(c.get("woonoppervlak_max")),
        "Inkomen": {"vereist": "laag inkomen verplicht", "bonus": "hoger bedrag bij laag inkomen"}.get(c.get("inkomen"), _tekst(c.get("inkomen"))),
        "Inkomensgrens": _tekst(r.get("inkomensgrens")),
        "VvE / appartement": " ".join(x for x in [_tekst(c.get("vve")), _tekst(c.get("vve_opmerking"))] if x),
        "Maatregelen": _tekst(r.get("maatregelen")), "Technische eisen": _tekst(r.get("technische_eisen")),
        "Eisen uitvoerder": _tekst(r.get("uitvoerder_eisen")), "Aanvragen": _tekst(r.get("aanvragen")),
        "Stapelbaar met ISDE": _tekst(r.get("stapelbaar_isde")), "Budget": _tekst(r.get("budgetstatus")),
        "Uitvoerder": _tekst(r.get("uitvoerder")), "Bewijs gecontroleerd": bewijs,
        "Opmerkingen": _tekst(r.get("opmerkingen")), "Peildatum": _datum(r.get("peildatum")),
        "Link": r.get("bron_url") or "", "Ook vermeld op": _tekst(r.get("extra_bronnen")),
    }


def schrijf_overzicht(regelingen, zoekstatus, dekking):
    """overzicht_voorwaarden.xlsx: blad Regelingen (één rij per regeling) en blad Gemeenten (alle gemeenten).
    Alleen om te lezen: wijzigingen in dit bestand worden niet teruggelezen (dat gaat via regelingen.json)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl ontbreekt: overzicht_voorwaarden.xlsx niet gemaakt (pip install openpyxl)")
        return
    kop_stijl, kop_vul = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5D50")
    geel = PatternFill("solid", fgColor="FFF2CC")

    def blad(ws, rijen, breedtes, link_kolom=None):
        koppen = list(rijen[0].keys()) if rijen else []
        ws.append(koppen)
        for cel in ws[1]:
            cel.font, cel.fill = kop_stijl, kop_vul
        for rij in rijen:
            ws.append([rij[k] for k in koppen])
        for i, k in enumerate(koppen, 1):
            ws.column_dimensions[get_column_letter(i)].width = breedtes.get(k, 18)
        for regel in ws.iter_rows(min_row=2):
            for cel in regel:
                cel.alignment = Alignment(wrap_text=True, vertical="top")
            if link_kolom and regel[koppen.index(link_kolom)].value:
                cel = regel[koppen.index(link_kolom)]
                cel.hyperlink, cel.font = cel.value, Font(color="0563C1", underline="single")
        ws.freeze_panes = "B2"
        if rijen:
            ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws = wb.active
    ws.title = "Regelingen"
    wachtend = [{"gemeente": z["gemeente"], "provincie": z.get("provincie"), "naam": w["titel"], "bron": "wacht",
                 "status": "nog niet uitgelezen", "bron_url": w["bron_url"],
                 "opmerkingen": "Gevonden in het CVDR, voorwaarden nog niet uitgelezen (volgt bij een volgende run)."}
                for z in zoekstatus for w in (z.get("wacht_op_uitlezen") or [])
                if not titel_valt_af(w["titel"], z["gemeente"])]
    rijen = [rij_voor(r) for r in sorted(regelingen + wachtend, key=lambda r: (r.get("gemeente") or "", r.get("naam") or ""))]
    blad(ws, rijen, {"Gemeente": 16, "Regeling": 40, "Bedrag": 35, "WOZ": 24, "Isolatiestaat / energielabel": 35,
                     "Maatregelen": 35, "Technische eisen": 35, "Eisen uitvoerder": 30, "Aanvragen": 30,
                     "Opmerkingen": 50, "Link": 45, "Inkomensgrens": 25, "VvE / appartement": 25,
                     "Bron": 24, "Bewijs gecontroleerd": 24, "Ook vermeld op": 45}, "Link")
    grijs = PatternFill("solid", fgColor="EDEDED")
    for regel, r in zip(ws.iter_rows(min_row=2), rijen):
        vul = grijs if r["Bron"] == "CVDR (wacht)" else geel if r["Bron"] not in ("CVDR", "handmatig") else None
        for cel in regel if vul else []:
            cel.fill = vul

    per_gemeente = {}
    for r in regelingen:
        per_gemeente.setdefault(r.get("gemeente"), []).append(r)
    status = {z["gemeente"]: z for z in zoekstatus}
    gemeenten = sorted(set(per_gemeente) | set(status) | set(dekking))
    g_rijen = []
    for g in gemeenten:
        regs, z, d = per_gemeente.get(g, []), status.get(g, {}), dekking.get(g, {})
        open_ = [r for r in regs if r.get("status") != "gesloten"]
        wacht = len(z.get("wacht_op_uitlezen") or [])
        bronnen = d.get("bronnen_gecheckt") or ["CVDR"]
        web_gedaan = "partners" in bronnen
        if open_:
            conclusie = "regeling gevonden"
        elif regs:
            conclusie = "alleen gesloten regeling(en)"
        elif wacht:
            conclusie = "wacht op uitlezen"
        elif web_gedaan and "gemeentesite" in bronnen:
            conclusie = "geen regeling gevonden (alles doorzocht)"
        elif web_gedaan:
            conclusie = "geen regeling gevonden (gemeentesite niet doorzocht)"
        else:
            conclusie = "nog niet volledig gezocht (alleen CVDR)"
        g_rijen.append({
            "Gemeente": g, "Provincie": z.get("provincie") or (regs[0].get("provincie") if regs else ""),
            "Conclusie": conclusie, "Regelingen (open)": len(open_), "Regelingen (totaal)": len(regs),
            "Waarvan via internet (AI)": sum(1 for r in regs if r.get("bron") == "web"),
            "Wacht op uitlezen": wacht, "Treffers in CVDR": d.get("treffers", ""),
            "CVDR gecheckt": _datum(z.get("cvdr_gecheckt") or d.get("datum")),
            "Bronnen doorzocht": ", ".join(bronnen),
            "Internet gezocht": _datum(z.get("internet_gezocht")),
            "Gemeentesite": z.get("gemeentesite") or ("" if not web_gedaan else "niet gevonden"),
            "Partners met regeling": ", ".join(z.get("partners") or []),
            "AI-zoekactie": _datum(z.get("ai_gezocht")),
            "Relevante pagina's": z.get("paginas_gevonden", ""),
            "Namen regelingen": ", ".join([r.get("naam") or "" for r in regs]
                                          + [f"{w['titel']} (nog niet uitgelezen)" for w in z.get("wacht_op_uitlezen") or []]),
        })
    blad(wb.create_sheet("Gemeenten"), g_rijen, {"Gemeente": 20, "Conclusie": 34, "Namen regelingen": 80,
                                                 "Bronnen doorzocht": 34, "Gemeentesite": 30, "Partners met regeling": 28})

    uitleg = wb.create_sheet("Uitleg")
    for regel in [
        ["Overzicht isolatieregelingen", f"bijgewerkt {_datum(VANDAAG)}"],
        ["Regelingen", "Eén rij per regeling met de voorwaarden. Geel = gevonden op internet (gemeentesite, partner of "
                       "AI-zoekactie), niet in het CVDR: altijd de link controleren. Grijs = gevonden in het CVDR maar nog niet uitgelezen. "
                       "'Bewijs gecontroleerd' = staat wat het model uitlas letterlijk in de brontekst. 'Ook vermeld op' = "
                       "webpagina's (bijv. van een energieloket) die dezelfde regeling noemen."],
        ["Gemeenten", "Alle gemeenten met de conclusie, ook als er niets is gevonden. 'Bronnen doorzocht' laat zien waar "
                      "is gezocht: CVDR, gemeentesite, partners (energieloketten e.d.) en AI-zoekactie. "
                      "'Nog niet volledig gezocht' betekent: alleen het CVDR, de rest volgt bij een volgende run."],
        ["Let op", "Dit bestand wordt bij elke run opnieuw gemaakt. Aanpassingen hierin gaan verloren; "
                   "'gecontroleerd' zet je in regelingen.json."],
    ]:
        uitleg.append(regel)
    uitleg.column_dimensions["A"].width, uitleg.column_dimensions["B"].width = 18, 110
    uitleg["A1"].font = Font(bold=True, size=13)
    wb.save(OVERZICHT)
    print(f"Overzicht geschreven: {OVERZICHT.name} ({len(rijen)} regelingen, {len(g_rijen)} gemeenten)")


def vergelijk(baseline_pad, resultaat):
    """Zet de modeluitkomst naast de handmatige controle (pilot)."""
    base = {r["cvdr_id"]: r for r in lees(ROOT / baseline_pad, []) if r.get("cvdr_id")}
    res = {r["cvdr_id"]: r for r in resultaat if r.get("cvdr_id")}
    regels = [f"# Vergelijking model vs. handmatig ({VANDAAG})\n"]
    for cid, b in base.items():
        r = res.get(cid)
        if not r or (r.get("opmerkingen") or "").startswith("Niet meer gevonden"):
            regels.append(f"## {b['gemeente']} – {cid}\nNIET GEVONDEN door de scraper.\n")
            continue
        d = verschillen(b, r)
        regels.append(f"## {b['gemeente']} – {cid}\n" + ("Alles gelijk.\n" if not d else
                      "".join(f"- {v}: handmatig {json.dumps(a, ensure_ascii=False)} | model {json.dumps(m, ensure_ascii=False)}\n" for v, a, m in d)))
    extra = [c for c in res if c not in base]
    if extra:
        regels.append("## Extra gevonden door de scraper\n" + "".join(f"- {res[c]['gemeente']}: {res[c]['naam']} ({c})\n" for c in extra))
    (DATA / "vergelijking.md").write_text("\n".join(regels), encoding="utf-8")
    print("Vergelijking geschreven naar data/vergelijking.md")


if __name__ == "__main__":
    main()
